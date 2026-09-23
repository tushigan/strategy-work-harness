"""Read back every exported cell and validation; rendering is a separate evidence item."""
import math
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from phase2_store import WorkflowError, current, events, local, sha, ref
from template_profile import DATE_FIELDS, SHEET, choices, matrix
from workbook_validations import expected_snapshot, validation_snapshot


def value_equal(actual, expected, field):
    if expected in ("", None):
        return actual in ("", None)
    if field in DATE_FIELDS and isinstance(actual, (datetime, date)):
        return actual.isoformat()[:10] == expected
    if field in DATE_FIELDS:
        return False
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and math.isclose(actual, expected)
    return actual == expected


def readback(path, profile, rows):
    book = openpyxl.load_workbook(path, data_only=False)
    errors = []
    try:
        if book.sheetnames != profile["sheets"]:
            errors.append("工作表名称或顺序改变")
        if SHEET not in book:
            return {"valid": False, "errors": errors + ["主表缺失"]}
        sheet = book[SHEET]
        headers = [c.value for c in sheet[1]]
        if headers != profile["headers"]:
            errors.append("导出表头或列顺序不一致")
        comments = {str(c.value): c.comment.text for c in sheet[1] if c.comment}
        if comments != profile["comments"]:
            errors.append("表头批注改变或丢失")
        for name, key in (("数据字典", "dictionary"), ("填写说明", "instructions")):
            if name not in book or matrix(book[name]) != profile[key]:
                errors.append(f"{name} 内容改变")
        for number, row in enumerate(rows, 2):
            for index, field in enumerate(profile["headers"], 1):
                cell = sheet.cell(number, index)
                if cell.data_type == "f" or not value_equal(cell.value, row.get(field), field):
                    errors.append(f"{cell.coordinate}: 回读值或类型与任务不符")
        for cells in sheet.iter_rows(min_row=len(rows) + 2):
            if any(c.value is not None for c in cells):
                errors.append(f"第{cells[0].row}行有多余数据/示例")
        native = profile.get("native_validations")
        if not native or validation_snapshot(path, SHEET) != expected_snapshot(native, len(rows)):
            errors.append("数据校验完整属性、提示、公式、联合范围或集合规则改变")
        for field, options in profile["options"].items():
            column = profile["headers"].index(field) + 1
            relevant = []
            for validation in sheet.data_validations.dataValidation:
                if any(part.min_col <= column <= part.max_col for part in validation.sqref.ranges):
                    if validation.type != "list" or choices(book, validation) != options:
                        errors.append(f"{field}: 数据校验选项改变")
                    else:
                        allowed_blank = {v["allow_blank"] for v in profile["validations"] if v["field"] == field}
                        if validation.allow_blank not in allowed_blank:
                            errors.append(f"{field}: 数据校验空值规则改变")
                        relevant.append(validation)
            for row_number in range(2, max(1001, len(rows) + 1) + 1):
                address = f"{get_column_letter(column)}{row_number}"
                if not any(address in validation for validation in relevant):
                    errors.append(f"{field}: 数据校验未覆盖 {address}")
                    break
        return {"valid": not errors, "errors": errors, "row_count": len(rows),
                "column_count": len(headers), "sheet_count": len(book.sheetnames)}
    finally:
        book.close()


def file_gate(root, meta=None):
    meta = meta or current(root, "gantt")
    checks = [e for e in events(root, "export-checks") if e.get("target") == ref(meta)]
    if not checks or not checks[-1].get("readback", {}).get("valid"):
        raise WorkflowError("当前 Excel 未完成实际文件回读")
    check = checks[-1]
    from compare_import_template import binding
    from gantt_model import ADAPTER_VERSION
    from validate_gantt import validate_current
    template = binding(root)
    if (check.get("template_sha256") != template["source"]["sha256"]
            or check.get("adapter_version") != ADAPTER_VERSION
            or check.get("plan_ref") != ref(current(root, "plan"))):
        raise WorkflowError("Excel 对应的模板、计划或适配器已改变，需生成新版本")
    plan_check = validate_current(root)
    if not plan_check["valid"]:
        raise WorkflowError("当前计划仍有问题，Excel 不能通过")
    actual = readback(local(root, meta["path"]), template["profile"], plan_check["rows"])
    if not actual["valid"]:
        raise WorkflowError("当前 Excel 重新回读失败")
    if not check.get("previews"):
        raise WorkflowError("当前 Excel 没有显示效果检核素材")
    for preview in check["previews"]:
        path = local(root, preview["path"])
        if not path.is_file() or sha(path) != preview["sha256"]:
            raise WorkflowError("Excel 显示效果素材丢失或改变")
    return check
