"""Preserve native validation metadata that the workbook writer cannot round-trip."""
import copy
import posixpath
import tempfile
import zipfile
from pathlib import Path
from xml.dom import minidom
from xml.etree import ElementTree as ET

from openpyxl.utils import get_column_letter, range_boundaries

from phase2_store import WorkflowError

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
VALIDATIONS = f"{{{MAIN}}}dataValidations"


def sheet_part(package, name):
    book = ET.fromstring(package.read("xl/workbook.xml"))
    sheet = next((s for s in book.findall(f"{{{MAIN}}}sheets/{{{MAIN}}}sheet")
                  if s.get("name") == name), None)
    if sheet is None:
        raise WorkflowError(f"工作簿缺工作表: {name}")
    relations = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    relation = next((r for r in relations.findall(f"{{{PACKAGE_REL}}}Relationship")
                     if r.get("Id") == sheet.get(f"{{{REL}}}id")), None)
    if relation is None or relation.get("TargetMode") == "External":
        raise WorkflowError("工作表关系缺失或指向外部文件")
    target = relation.get("Target", "")
    part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else
                             posixpath.join("xl", target))
    if not part.startswith("xl/") or part not in package.namelist():
        raise WorkflowError("工作表关系路径无效")
    return part


def snapshot_node(node):
    attrs = dict(node.attrib)
    if "sqref" in attrs:
        attrs["sqref"] = " ".join(sorted(attrs["sqref"].split()))
    return {"tag": node.tag, "attrs": attrs,
            "text": node.text if not len(node) else None,
            "children": [snapshot_node(child) for child in node]}


def validation_snapshot(path, sheet):
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read(sheet_part(package, sheet)))
        nodes = root.findall(VALIDATIONS)
        if len(nodes) > 1:
            raise WorkflowError("同工作表存在多个数据校验集合，需人工适配")
        return snapshot_node(nodes[0]) if nodes else None


def extended_range(value, last_row):
    parts = []
    for part in value.split():
        left, top, right, bottom = range_boundaries(part)
        if top != 2 or bottom != 1001 or None in (left, right):
            raise WorkflowError("非基线数据校验范围不能自动延展，需重新适配")
        parts.append(f"{get_column_letter(left)}{top}:{get_column_letter(right)}{max(bottom, last_row)}")
    return " ".join(sorted(parts))


def expected_snapshot(snapshot, row_count):
    result = copy.deepcopy(snapshot)
    if result:
        for rule in result["children"]:
            rule["attrs"]["sqref"] = extended_range(rule["attrs"]["sqref"], row_count + 1)
    return result


def adaptation_signature(profile):
    result = copy.deepcopy(profile["native_validations"])
    if result:
        for rule in result["children"]:
            parts = []
            for part in rule["attrs"]["sqref"].split():
                left, top, right, bottom = range_boundaries(part)
                fields = profile["headers"][left - 1:right]
                parts.append({"fields": fields, "start": top, "end": bottom})
            rule["attrs"]["sqref"] = sorted(parts, key=str)
    return result


def restore_validations(template, output, sheet, row_count):
    """Change only the validation subtree in a newly authored temporary workbook."""
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve():
        raise WorkflowError("不得用导出覆盖源模板")
    with zipfile.ZipFile(template) as source, zipfile.ZipFile(output) as target:
        source_doc = minidom.parseString(source.read(sheet_part(source, sheet)))
        part = sheet_part(target, sheet)
        output_doc = minidom.parseString(target.read(part))
        try:
            original = source_doc.getElementsByTagNameNS(MAIN, "dataValidations")
            existing = output_doc.getElementsByTagNameNS(MAIN, "dataValidations")
            if len(original) != 1 or len(existing) != 1:
                raise WorkflowError("缺失或重复的数据校验集合，停止导出")
            replacement = output_doc.importNode(original[0], deep=True)
            # Preserve namespace declarations when copying across worksheet documents.
            for name, value in source_doc.documentElement.attributes.items():
                if name == "xmlns" or name.startswith("xmlns:"):
                    replacement.setAttribute(name, value)
            for rule in replacement.getElementsByTagNameNS(MAIN, "dataValidation"):
                rule.setAttribute("sqref", extended_range(rule.getAttribute("sqref"), row_count + 1))
            output_doc.documentElement.replaceChild(replacement, existing[0])
            data = output_doc.toxml(encoding="utf-8")
            with tempfile.TemporaryDirectory(prefix="harness-validation-", dir=output.parent) as folder:
                patched = Path(folder) / "patched.xlsx"
                with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as package:
                    for item in target.infolist():
                        package.writestr(item, data if item.filename == part else target.read(item.filename))
                # Close the input before replacing it, including on Windows.
                target.close()
                patched.replace(output)
        finally:
            source_doc.unlink()
            output_doc.unlink()
