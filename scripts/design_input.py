"""Decode local design formats and reuse the existing design-brief schema."""
from pathlib import Path

from phase2_store import WorkflowError
from phase3_document import DESIGN_FIELDS, DESIGN_TYPES


def pages(path):
    try:
        return _pages(path)
    except Exception as exc:
        raise WorkflowError(f"设计稿读取失败：{exc}") from exc


def _pages(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size == 0:
        raise WorkflowError("设计稿文件为空或不存在")
    suffix = path.suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image
        with Image.open(path) as image:
            if image.format != ("PNG" if suffix == ".png" else "JPEG"):
                raise WorkflowError("图片内容与后缀不符")
            image.load()
        return 1, "image"
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted or not reader.pages:
            raise WorkflowError("PDF加密或没有可读页面")
        for page in reader.pages:
            page.get_contents()
        return len(reader.pages), "pdf"
    if suffix == ".pptx":
        from pptx import Presentation
        deck = Presentation(str(path))
        if not deck.slides:
            raise WorkflowError("PPTX没有可读幻灯片")
        return len(deck.slides), "pptx"
    raise WorkflowError("支持 PNG、JPG、JPEG、PDF、PPTX；旧式 PPT 需先转换")


def strategy_checks(data):
    design = data.get("design", {})
    missing = [name for name in DESIGN_FIELDS if not design.get(name)]
    if data.get("brief_type") not in DESIGN_TYPES:
        missing.append("design brief_type")
    checks = []
    for number, raw in enumerate(design.get("observable_checks", []), 1):
        if not isinstance(raw, dict) or not all(isinstance(raw.get(k), str) and raw[k].strip() for k in (
                "question", "basis")) or (
                raw.get("level") not in {"red", "yellow"}):
            missing.append(f"observable_checks {number}")
            continue
        checks.append({**raw, "id": f"CHECK-{number:03d}"})
    if not checks:
        missing.append("observable_checks")
    return checks, missing


def validate_report_coverage(root, item, report):
    from phase5_visual import image_evidence
    passing = report["status"] in {"passed", "passed_with_yellow"}
    expected = set(range(1, item["pages"] + 1))
    checked = report.get("checked_pages")
    visual = report.get("visual_evidence")
    if not isinstance(checked, list) or any(type(p) is not int or p not in expected for p in checked):
        raise WorkflowError("checked_pages 必须是实际有效页码列表")
    if not isinstance(visual, list):
        raise WorkflowError("必须提供逐页视觉证据")
    seen = set()
    evidence_owners = {}
    for entry in visual:
        if not isinstance(entry, dict) or type(entry.get("page")) is not int or entry["page"] not in expected:
            raise WorkflowError("视觉证据页码无效")
        if entry.get("rendered") is not True:
            continue
        if not all(isinstance(entry.get(k), str) and entry[k].strip() for k in ("method", "observation")):
            raise WorkflowError("视觉证据需要打开方法和实质画面观察")
        image_evidence(root, entry.get("file"))
        from phase5_visual import evidence_identity
        identity = evidence_identity(entry["file"])
        owner = evidence_owners.get(identity)
        if owner is not None and owner != entry["page"]:
            raise WorkflowError("不同页面不能复用同一张截图证据")
        evidence_owners[identity] = entry["page"]
        seen.add(entry["page"])
    if passing and (item["parse_error"] or not expected or set(checked) != expected or seen != expected or not (
            {f"page {p}" for p in expected}.issubset(report["scope"]))):
        raise WorkflowError("不可读或未逐页完整检核，不能通过")
    if not passing and (seen != expected or item["parse_error"]) and not report["uncovered"]:
        raise WorkflowError("读取失败或未完整检核必须保留未覆盖范围")
    checks = {x["id"]: x for x in item["strategy_checks"]}
    for finding in report["findings"]:
        if type(finding.get("page")) is not int or finding["page"] not in expected or not finding.get("region"):
            raise WorkflowError("每条设计意见需要实际页码和具体区域")
        basis = checks.get(finding.get("basis_id"))
        if not basis:
            raise WorkflowError("设计意见必须引用简报中的检查项编号")
        if finding["level"] == "red" and basis["level"] != "red":
            raise WorkflowError("可讨论的检查项不能因审美差异升级为红灯")
    if passing and not set(checks).issubset({f.get("basis_id") for f in report["findings"]}):
        raise WorkflowError("通过报告没有覆盖全部策略检查项")
