"""Read template rules without rewriting the supplied workbook."""
from __future__ import annotations

from datetime import date, datetime
import openpyxl

from phase2_store import WorkflowError, fingerprint, sha
from workbook_validations import validation_snapshot

BASELINE_SHA = "ae3c8d041ff8aa9a2fcc48a7ceb32d2c0b7e54950b4a307a132b7e19f5a0d59d"
SHEET = "批量立项模板"
DATE_FIELDS = ["合同签约时间", "合同完成时间", "任务开始", "任务截止"]
PEOPLE_FIELDS = ["合同主项目经理", "合同助理项目经理", "执行者（可多人）", "策划负责人", "设计负责人", "策略总监", "创意总监", "最终拍板人"]
TASK_FIELDS = ["任务", "任务类型", "任务状态", "任务开始", "任务开始半天", "任务截止", "任务截止半天", "执行者（可多人）", "任务等级", "交付物类型", "验收标准", "任务说明", "策划负责人", "设计负责人", "策略总监", "创意总监", "最终拍板人"]


def scalar(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def matrix(sheet):
    return [[scalar(c.value) for c in row] for row in sheet.iter_rows()]


def choices(book, validation):
    formula = (validation.formula1 or "").lstrip("=")
    if formula.startswith('"') and formula.endswith('"'):
        return formula[1:-1].split(",")
    if "!" in formula:
        tab, address = formula.rsplit("!", 1)
        tab = tab.strip("'").replace("''", "'")
        if tab in book.sheetnames:
            try:
                return [str(c.value) for row in book[tab][address] for c in row if c.value is not None]
            except (ValueError, TypeError, AttributeError, KeyError) as exc:
                raise WorkflowError(f"下拉范围无法解析: {formula!r}") from exc
    raise WorkflowError(f"暂不支持下拉公式 {formula!r}，需要人工核对适配")


def inspect_template(path):
    # 只读分析不等于 openpyxl 的 read_only 模式：该模式会隐藏数据校验对象。
    # 这里不保存、不修改工作簿，只用普通加载读取规则快照。
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        issues = [f"模板缺工作表: {name}" for name in (SHEET, "数据字典", "填写说明")
                  if name not in book.sheetnames]
        if SHEET not in book.sheetnames:
            return {"sha256": sha(path), "semantic_sha256": fingerprint(book.sheetnames),
                    "sheets": book.sheetnames, "headers": [], "rules": {}, "comments": {},
                    "options": {}, "dictionary": [], "instructions": [], "validations": [],
                    "native_validations": None, "sample_rows": [], "old_data_rows": 0, "issues": issues}
        sheet = book[SHEET]
        headers = [c.value for c in sheet[1]]
        if not all(isinstance(h, str) and h.strip() for h in headers) or len(set(headers)) != len(headers):
            issues.append("模板表头为空或重复，不能自动映射")
        rules = {}
        instructions = matrix(book["填写说明"]) if "填写说明" in book.sheetnames else []
        for row in instructions:
            if row and row[0] in headers and isinstance(row[0], str):
                padded = row + [None] * 6
                rules[row[0]] = {"required": padded[1], "allowed": padded[2],
                                "default": padded[3], "description": padded[4], "example": padded[5]}
        if set(rules) != set(headers):
            issues.append("模板有字段没有填写规则，请补充适配后再生成")
        validations = []
        options = {}
        for item in sheet.data_validations.dataValidation:
            if item.type != "list":
                issues.append("出现新的校验类型，需要适配")
                continue
            try:
                values = choices(book, item)
            except WorkflowError as exc:
                issues.append(str(exc))
                continue
            for part in sorted(item.sqref.ranges, key=str):
                for index in range(part.min_col - 1, part.max_col):
                    if index >= len(headers):
                        issues.append("数据校验越过已知表头")
                        continue
                    field = headers[index]
                    if field in options and options[field] != values:
                        issues.append(f"同字段下拉规则冲突: {field}")
                    options[field] = values
                    validations.append({"field": field, "column": index, "values": values,
                                        "range": str(part), "allow_blank": item.allow_blank})
        validations.sort(key=lambda v: (v["column"], v["range"], str(v["values"])))
        comments = {str(c.value): c.comment.text for c in sheet[1] if c.comment}
        semantic = {"sheets": book.sheetnames, "headers": headers, "rules": rules,
                    "comments": comments, "options": options,
                    "dictionary": matrix(book["数据字典"]) if "数据字典" in book.sheetnames else [],
                    "instructions": instructions,
                    "validations": validations,
                    "native_validations": validation_snapshot(path, SHEET)}
        return {"sha256": sha(path), "semantic_sha256": fingerprint(semantic), **semantic,
                "sample_rows": matrix(sheet)[1:], "old_data_rows": max(1, sheet.max_row - 1),
                "issues": issues}
    finally:
        book.close()


def compare_profiles(old, new):
    changes = []
    for key in ("sheets", "headers", "rules", "comments", "options", "dictionary",
                "instructions", "validations", "native_validations"):
        if old[key] != new[key]:
            changes.append({"section": key, "before": old[key], "after": new[key]})
    return {"file_changed": old["sha256"] != new["sha256"],
            "rules_changed": bool(changes), "changes": changes,
            "rules_unchanged": not changes,
            "requires_adaptation": bool(changes)}


def required_fields(profile):
    return [field for field, rule in profile["rules"].items() if "必填" in str(rule["required"])]


def default_effects(profile, row):
    return [{"field": field, "default": rule["default"], "rule": rule["description"]}
            for field, rule in profile["rules"].items()
            if rule["default"] not in (None, "", "无")
            and (row.get(field) in (None, "", []) or isinstance(row.get(field), str) and not row[field].strip())]
