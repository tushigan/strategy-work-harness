"""Explicit, version-bound transcription checks; never a whole-source bypass."""
from pathlib import Path
import re

from phase2_store import WorkflowError, archive, local, now, read_json, sha
from workspace_lock import serialized

DETAILS = ("verified_by", "recorded_by", "reason", "verification_action")


def scope_for(root, original, data, location):
    if not isinstance(location, str):
        raise WorkflowError("核验须指定一页 PDF 或一个单元格范围")
    suffix = Path(original["path"]).suffix.lower()
    if suffix == ".pdf":
        found = re.fullmatch(r"第([1-9]\d*)页", location)
        if not found or int(found[1]) not in {u.get("page") for u in data["units"]}:
            raise WorkflowError("PDF 核验须逐页指定实际存在的页码，例如第1页")
        return {"kind": "pdf_page", "page": int(found[1])}, location
    if suffix not in {".xlsx", ".xlsm"}:
        raise WorkflowError("按范围补录核验仅支持已登记的 PDF/XLSX/XLSM，不代替音频转写或 DOCX 读取")
    found = re.fullmatch(r"'((?:[^']|'')+)'!([A-Za-z]{1,3}[1-9]\d*)(?::([A-Za-z]{1,3}[1-9]\d*))?", location)
    if not found:
        raise WorkflowError("单元格核验需明确工作表与范围，例如 'Sheet'!A2:B2，不接受整表通配符")
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
    sheet = found[1].replace("''", "'")
    cells = f"{found[2].upper()}:{(found[3] or found[2]).upper()}"
    left, top, right, bottom = range_boundaries(cells)
    if left > right or top > bottom or right > 16384 or bottom > 1048576:
        raise WorkflowError("核验单元格范围反向或越界")
    book = openpyxl.load_workbook(local(root, original["path"]), read_only=True, data_only=False)
    try:
        if sheet not in book.sheetnames:
            raise WorkflowError("核验工作表不存在于归档原件")
        tab = book[sheet]
        if right > tab.max_column or bottom > tab.max_row:
            raise WorkflowError("核验单元格范围超出原件已有区域")
    finally:
        book.close()
    scope = {"kind": "xlsx_range", "sheet": sheet, "cells": cells,
             "bounds": [left, top, right, bottom]}
    escaped_sheet = sheet.replace("'", "''")
    return scope, f"'{escaped_sheet}'!{cells}"


def overlaps(first, second):
    if first["kind"] != second["kind"]:
        return False
    if first["kind"] == "pdf_page":
        return first["page"] == second["page"]
    a, b = first["bounds"], second["bounds"]
    return first["sheet"] == second["sheet"] and not (
        a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1])


def verification_files(data):
    return [record[field] for record in data.get("verified_ranges", [])
            for field in ("base_file", "transcript")]


def validate_verifications(root, meta, data, history):
    from phase3_sources import file_valid, source_ref
    records = data["verified_ranges"]
    if not isinstance(records, list) or not records:
        raise WorkflowError("资料核验记录不完整，不能按整份已核验放行")
    scopes = []
    for record in records:
        if not isinstance(record, dict) or any(
                not isinstance(record.get(k), str) or not record[k].strip()
                for k in (*DETAILS, "location", "text", "recorded_at")):
            raise WorkflowError("资料核验缺范围、转录、核验人、理由或实际动作")
        base_ref = record.get("base_source")
        base = next((m for m in history if source_ref(m) == base_ref), None)
        if not base or base["version"] >= meta["version"] or (
                type(base_ref.get("version")) is not int) or base["original"] != meta["original"]:
            raise WorkflowError("核验绑定的原件或基础数据版本不符")
        if record.get("base_file") != {"path": base["path"], "sha256": base["sha256"]} or (
                record.get("original_sha256") != meta["original"]["sha256"]):
            raise WorkflowError("核验绑定的原件或基础数据指纹不符")
        if any(not file_valid(root, record.get(field)) for field in ("base_file", "transcript")):
            raise WorkflowError("核验基础数据或转录改变/丢失，需重新核验")
        prior = read_json(local(root, base["path"]))
        raw_prior = {k: v for k, v in prior.items() if k != "verified_ranges"}
        raw_data = {k: v for k, v in data.items() if k != "verified_ranges"}
        if raw_prior != raw_data:
            raise WorkflowError("核验不能改变原件提取记录或完整性标记")
        if local(root, record["transcript"]["path"]).read_text(encoding="utf-8-sig") != record["text"]:
            raise WorkflowError("核验转录与归档原文不符")
        scope, location = scope_for(root, meta["original"], data, record["location"])
        if record.get("scope") != scope or record["location"] != location or any(
                overlaps(scope, other) for other in scopes):
            raise WorkflowError("核验范围不符或相互重叠，不能扩大使用")
        scopes.append(scope)


def inspect_source(root, source_id):
    from phase3_sources import source_current, source_ref
    meta = source_current(root, source_id)
    data = read_json(local(root, meta["path"]))
    records = data.get("verified_ranges")
    usable = records if records is not None else data["units"] if data["extraction_complete"] else []
    if data["metadata"]["relation"] == "method_reference":
        usable = []
    return {"source": source_ref(meta), "path": meta["path"], "original": meta["original"],
            "data": data, "usable_locations": [u["location"] for u in usable],
            "scope_notice": "核验只覆盖列出范围；不证明原文陈述真实或市场认可，不代替独立检核"}


def verify_source(root, reference, original_sha256, location, transcript_path,
                  verified_by, reason, verification_action, actor):
    from phase3_sources import _save_source, catalog, source_current, source_ref
    with serialized(root):
        if not isinstance(reference, dict) or type(reference.get("version")) is not int:
            raise WorkflowError("资料核验须绑定当前来源版本和数据指纹")
        meta = source_current(root, reference.get("key"))
        if reference != source_ref(meta):
            raise WorkflowError("资料核验请求不是当前来源版本/数据指纹；先回读 source-status")
        if original_sha256 != meta["original"]["sha256"]:
            raise WorkflowError("资料核验的原件指纹不符")
        details = dict(zip(DETAILS, (verified_by, actor, reason, verification_action)))
        if any(not isinstance(value, str) or not value.strip() for value in details.values()):
            raise WorkflowError("核验须提供实际核验人、登记者、理由和已执行的核验动作")
        if transcript_path is None:
            raise WorkflowError("请提供本范围的 UTF-8 TXT/Markdown 转录文件，不需手写 JSON")
        transcript_path = Path(transcript_path)
        if transcript_path.suffix.lower() not in {".txt", ".md"} or not transcript_path.is_file():
            raise WorkflowError("转录须为存在的 UTF-8 TXT/Markdown 文件")
        text = transcript_path.read_text(encoding="utf-8-sig")
        if not text.strip():
            raise WorkflowError("核验转录不能为空，不能用单个确认标记放行整份资料")
        data = read_json(local(root, meta["path"]))
        scope, location = scope_for(root, meta["original"], data, location)
        records = data.get("verified_ranges", [])
        same = next((r for r in records if r["location"] == location and r["text"] == text
                     and all(r.get(k) == v for k, v in details.items())
                     and r["transcript"]["sha256"] == sha(transcript_path)), None)
        if same:
            return meta
        transcript = archive(root, transcript_path, "evidence")
        if local(root, transcript["path"]).read_text(encoding="utf-8-sig") != text or (
                source_ref(source_current(root, meta["source_id"])) != reference):
            raise WorkflowError("归档期间来源或转录改变，重新回读后核验")
        record = {"location": location, "scope": scope, "text": text, "transcript": transcript,
                  "base_source": reference, "base_file": {"path": meta["path"], "sha256": meta["sha256"]},
                  "original_sha256": original_sha256, "recorded_at": now(), **details}
        # A partial replacement cannot keep the wider old range implicitly approved.
        data["verified_ranges"] = [r for r in records if not overlaps(r["scope"], scope)] + [record]
        result = _save_source(root, data, actor, catalog(root))
        source_current(root, result["source_id"])
        return result
